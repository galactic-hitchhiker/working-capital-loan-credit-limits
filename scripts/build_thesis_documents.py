#!/usr/bin/env python3
"""Build the formatted thesis and committee report Word documents.

The Markdown files in docs/thesis remain the editable source.  Pandoc is used
for native Word equations and tables; python-docx applies the RAF reference
layout and generates the formal report.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
THESIS_DIR = ROOT / "docs" / "thesis"
OUTPUT_DIR = ROOT / "docs" / "deliverables"
BUILD_DIR = ROOT / "build" / "thesis_documents"
REFERENCE_DIR = Path("/private/tmp/wcl-thesis-refs.5Cpxri")
THESIS_TEMPLATE = REFERENCE_DIR / "thesis_example.docx"
REPORT_TEMPLATE = REFERENCE_DIR / "report_example.docx"

THESIS_DOCX = OUTPUT_DIR / "Djordje_Velickovic_Master_Thesis.docx"
THESIS_PDF = OUTPUT_DIR / "Djordje_Velickovic_Master_Thesis.pdf"
REPORT_DOCX = OUTPUT_DIR / "Djordje_Velickovic_Master_Thesis_Report.docx"
REPORT_PDF = OUTPUT_DIR / "Djordje_Velickovic_Master_Thesis_Report.pdf"
BASE_FONT_SIZE = 12

BODY_FILES = [
    "02_chapter_1.md",
    "03_chapter_2.md",
    "04_chapter_3.md",
    "05_chapter_4.md",
    "06_chapter_5.md",
    "07_chapter_6.md",
    "08_chapter_7.md",
    "09_chapter_8.md",
    "10_references.md",
    "11_appendix_a_reproducibility.md",
]


def run(*args: str) -> None:
    subprocess.run(args, check=True)


def normalize_math(markdown: str) -> str:
    """Convert the source's inline $$...$$ convention to Pandoc inline math.

    Standalone equations stay display math. Equation tags remain visible for
    Pandoc and are repositioned at the right margin during DOCX formatting.
    """
    output: list[str] = []
    in_display = False
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped == "$$":
            in_display = not in_display
            output.append(line)
            continue
        if in_display or (stripped.startswith("$$") and stripped.endswith("$$")):
            line = re.sub(r"\\tag\{(\d+)\}", r"\\qquad (\1)", line)
            output.append(line)
            continue
        line = re.sub(r"\$\$(.+?)\$\$", r"$\1$", line)
        # A standalone Markdown image treats its alt text as an automatic
        # caption in DOCX. The thesis already has explicit numbered captions.
        line = re.sub(r"!\[[^\]]*\]\((figures/[^)]+)\)", r"![](\1)", line)
        output.append(line)
    return "\n".join(output) + "\n"


def parse_headings() -> list[tuple[int, str]]:
    headings: list[tuple[int, str]] = [(1, "Abstract")]
    for filename in BODY_FILES:
        for line in (THESIS_DIR / filename).read_text(encoding="utf-8").splitlines():
            match = re.match(r"^(#{1,3})\s+(.+?)\s*$", line)
            if match:
                headings.append((len(match.group(1)), match.group(2)))
    return headings


def parse_captions() -> tuple[list[str], list[str]]:
    figures: list[str] = []
    tables: list[str] = []
    for filename in BODY_FILES:
        text = (THESIS_DIR / filename).read_text(encoding="utf-8")
        for kind, number, caption in re.findall(
            r"^\*(Figure|Table)\s+(\d+)\.\s+(.+?)\*\s*$", text, re.MULTILINE
        ):
            full = f"{kind} {number}. {caption}"
            (figures if kind == "Figure" else tables).append(full)
    figures.sort(key=lambda value: int(re.match(r"Figure (\d+)", value).group(1)))
    tables.sort(key=lambda value: int(re.match(r"Table (\d+)", value).group(1)))
    return figures, tables


def marker_entry(prefix: str, level: int, title: str, page: int) -> str:
    return f"@{prefix}{level}@ {title} ¦ {page}"


def make_markdown(page_map: dict[str, int]) -> str:
    headings = parse_headings()
    figures, tables = parse_captions()
    abstract = normalize_math((THESIS_DIR / "01_abstract.md").read_text(encoding="utf-8"))

    logo = REFERENCE_DIR / "thesis_unzip" / "word" / "media" / "image1.jpg"
    cover = [
        f"![]({logo.as_posix()}){{width=2.4cm}}",
        "",
        "@COVER-UNIVERSITY@ UNION UNIVERSITY",
        "",
        "@COVER-SCHOOL@ SCHOOL OF COMPUTING BELGRADE, SERBIA",
        "",
        "@COVER-PROGRAM@ MASTER IN COMPUTATIONAL FINANCE",
        "",
        "@COVER-DOC@ MASTER'S THESIS",
        "",
        "@COVER-TITLE@ Dynamic Determination of Short-Term Corporate Credit Limits Using Machine Learning",
        "",
        "@COVER-STUDENT@ Student: Đorđe Veličković",
        "",
        "@COVER-ID@ Student ID number: [TO BE ENTERED]",
        "",
        "@COVER-SUPERVISOR@ Supervisor: Prof. Dr. Branko Urošević",
        "",
        "@COVER-PLACE@ Belgrade, October 2026",
        "",
        "@PAGEBREAK@",
        "",
    ]

    front: list[str] = [abstract.rstrip(), "", "@PAGEBREAK@", "", "# Table of Contents", ""]
    for level, title in headings:
        front.extend([marker_entry("TOC", level, title, page_map.get(title, 0)), ""])
    front.extend(["@PAGEBREAK@", "", "# List of Figures", ""])
    for caption in figures:
        front.extend([marker_entry("FIG", 1, caption, page_map.get(caption, 0)), ""])
    front.extend(["@PAGEBREAK@", "", "# List of Tables", ""])
    for caption in tables:
        front.extend([marker_entry("TAB", 1, caption, page_map.get(caption, 0)), ""])
    front.extend(["@PAGEBREAK@", ""])

    body: list[str] = []
    for index, filename in enumerate(BODY_FILES):
        if index:
            body.extend(["", "@PAGEBREAK@", ""])
        body.append(normalize_math((THESIS_DIR / filename).read_text(encoding="utf-8")).rstrip())
    return "\n".join(cover + front + body) + "\n"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, BASE_FONT_SIZE)


def right_align_equation_numbers(doc: Document) -> None:
    """Keep display equations centered and place their numbers at the right margin.

    Pandoc converts the LaTeX equation tag into part of the editable OMML equation. Split
    the trailing number from that math object and use Word tab stops, which is
    the conventional single-line layout for a centered equation with a
    right-aligned number.
    """

    for paragraph in doc.paragraphs:
        math_paragraphs = paragraph._p.xpath("./m:oMathPara")
        if len(math_paragraphs) != 1:
            continue
        math_paragraph = math_paragraphs[0]
        math = math_paragraph.find(qn("m:oMath"))
        if math is None:
            continue

        children = list(math)

        def math_text(element) -> str:
            return "".join(node.text or "" for node in element.iter(qn("m:t")))

        if len(children) < 3:
            continue
        tail = [math_text(element) for element in children[-3:]]
        if tail[0] != "(" or not re.fullmatch(r"\d+", tail[1]) or tail[2] != ")":
            continue
        equation_number = tail[1]
        for element in children[-3:]:
            math.remove(element)
        if len(math) and math_text(math[-1]).isspace():
            math.remove(math[-1])

        # Convert the block-math wrapper to an inline OMML object so tabs can
        # position the equation and its number independently on the same line.
        math_paragraph.remove(math)
        paragraph._p.remove(math_paragraph)
        prefix = paragraph.add_run("\t")
        set_run_font(prefix, BASE_FONT_SIZE)
        paragraph._p.append(math)
        number_run = paragraph.add_run(f"\t({equation_number})")
        set_run_font(number_run, BASE_FONT_SIZE)

        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = None
        paragraph.paragraph_format.right_indent = None
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.keep_together = True
        tabs = paragraph.paragraph_format.tab_stops
        tabs.clear_all()
        tabs.add_tab_stop(Cm(8.25), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)


def apply_thesis_format(path: Path) -> None:
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    styles = doc.styles
    for style_name in ("Normal", "Body Text", "Normal (Web)"):
        if style_name in styles:
            style = styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            style.font.size = Pt(BASE_FONT_SIZE)
            style.font.bold = False
            style.font.italic = False
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            style.paragraph_format.line_spacing = 1.15
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(10)
            style.paragraph_format.first_line_indent = None
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.left_indent = None
        style.paragraph_format.right_indent = None
        style.paragraph_format.first_line_indent = None
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)
    if "Caption" in styles:
        caption_style = styles["Caption"]
    else:
        caption_style = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption_style.font.name = "Times New Roman"
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption_style.font.size = Pt(BASE_FONT_SIZE)
    caption_style.font.italic = True
    caption_style.font.color.rgb = RGBColor(0, 0, 0)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Captions follow tables and figures, so do not attach them to the next
    # body paragraph. The object itself is kept with the caption below it.
    caption_style.paragraph_format.keep_with_next = False
    caption_style.paragraph_format.keep_together = True
    caption_style.paragraph_format.space_before = Pt(4)
    caption_style.paragraph_format.space_after = Pt(8)

    pending_page_break = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "@PAGEBREAK@":
            pending_page_break = True
            # Remove the marker paragraph completely. A cleared paragraph can
            # itself flow onto a new page and make the following break create
            # a second, blank page.
            p._element.getparent().remove(p._element)
            continue
        if pending_page_break and text:
            # Applying the break to the next real paragraph avoids an empty
            # page when the previous content already ended at a page boundary.
            p.paragraph_format.page_break_before = True
            pending_page_break = False
        cover_match = re.match(r"^@COVER-([A-Z]+)@\s*(.*)$", text)
        if cover_match:
            kind, value = cover_match.groups()
            p.clear()
            run = p.add_run(value)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            settings = {
                "UNIVERSITY": (14, True, 6, 0),
                "SCHOOL": (14, True, 0, 10),
                "PROGRAM": (13, True, 4, 24),
                "DOC": (16, True, 16, 16),
                "TITLE": (16, True, 6, 28),
                "STUDENT": (12, False, 6, 0),
                "ID": (12, False, 0, 0),
                "SUPERVISOR": (12, False, 0, 30),
                "PLACE": (12, False, 0, 0),
            }
            size, bold, before, after = settings[kind]
            set_run_font(run, size, bold=bold)
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.space_after = Pt(after)
            p.paragraph_format.line_spacing = 1.0
            continue
        marker = re.match(r"^@(TOC|FIG|TAB)([123])@\s*(.*?)\s*¦\s*(\d+)\s*$", text)
        if marker:
            kind, level_s, title, page = marker.groups()
            level = int(level_s)
            p.clear()
            title_run = p.add_run(title)
            p.add_run("\t")
            page_run = p.add_run(page)
            for run in (title_run, page_run):
                set_run_font(run, BASE_FONT_SIZE)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.55 * (level - 1))
            tabs = p.paragraph_format.tab_stops
            tabs.add_tab_stop(Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT, 1)
            continue
        if re.match(r"^(Table|Figure)\s+\d+\.\s+", text):
            p.style = styles["Caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = None
            p.paragraph_format.right_indent = None
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.tab_stops.clear_all()
        if p.style.name in ("Normal", "Body Text", "Normal (Web)") and text:
            is_list = bool(p._p.xpath("./w:pPr/w:numPr"))
            is_equation = bool(p._p.xpath(".//m:oMath"))
            p.paragraph_format.left_indent = None
            p.paragraph_format.right_indent = None
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(10)
            if is_equation:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
            elif is_list:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.first_line_indent = None
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.first_line_indent = None
        for run in p.runs:
            if not text or text.startswith("@COVER"):
                continue
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if p.style.name.startswith("Heading"):
                heading_size = {"Heading 1": 16, "Heading 2": 13, "Heading 3": 12}.get(
                    p.style.name, BASE_FONT_SIZE
                )
                run.font.size = Pt(heading_size)
                run.bold = True
            elif p.style.name in ("Normal", "Body Text", "Normal (Web)"):
                run.font.size = Pt(BASE_FONT_SIZE)
                run.bold = False
            elif run.font.size is None:
                run.font.size = Pt(BASE_FONT_SIZE)

    right_align_equation_numbers(doc)

    # Front-matter headings are centered; chapter headings keep the reference layout.
    for p in doc.paragraphs:
        if p.text in ("Abstract", "Table of Contents", "List of Figures", "List of Tables"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.text == "Abstract":
            p.runs[0].font.size = Pt(14)
            p.runs[0].bold = False

    for table_number, table in enumerate(doc.tables, start=1):
        table.style = styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.first_child_found_in("w:tblW")
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.insert(0, tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")
        tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
        if tbl_layout is None:
            tbl_layout = OxmlElement("w:tblLayout")
            tbl_pr.append(tbl_layout)
        tbl_layout.set(qn("w:type"), "fixed")
        column_count = len(table.columns)
        if column_count == 2:
            ratios = [0.74, 0.26]
        elif column_count == 3:
            ratios = [0.36, 0.34, 0.30]
        elif column_count == 4:
            ratios = [0.31, 0.31, 0.19, 0.19]
        elif column_count == 5:
            ratios = [0.28, 0.18, 0.18, 0.18, 0.18]
        elif column_count == 6:
            ratios = [0.25, 0.15, 0.15, 0.15, 0.15, 0.15]
        elif column_count == 7:
            ratios = [0.25, 0.125, 0.15, 0.10, 0.15, 0.125, 0.10]
        else:
            ratios = [1 / column_count] * column_count
        if table_number == 11:
            ratios = [0.30, 0.50, 0.20]
        total_twips = 8787
        widths = [int(total_twips * ratio) for ratio in ratios]
        for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_col.set(qn("w:w"), str(width))
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for r_index, row in enumerate(table.rows):
            set_row_no_split(row)
            for c_index, cell in enumerate(row.cells):
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.first_child_found_in("w:tcW")
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.insert(0, tc_w)
                tc_w.set(qn("w:type"), "dxa")
                tc_w.set(qn("w:w"), str(widths[min(c_index, len(widths) - 1)]))
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                for p in cell.paragraphs:
                    # Pandoc labels table-cell paragraphs as "Compact".  The
                    # supplied reference DOCX does not define that style, and
                    # LibreOffice then lays the cell text outside the grid.
                    p.style = styles["Table Paragraph"]
                    if p.alignment is None:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = None
                    if r_index == len(table.rows) - 1:
                        p.paragraph_format.keep_with_next = True
                    for run in p.runs:
                        set_run_font(run, BASE_FONT_SIZE, bold=(r_index == 0))

    # Constrain all inline figures to the printable width.
    for shape in doc.inline_shapes:
        if shape.width > Cm(15.5):
            ratio = Cm(15.5) / shape.width
            shape.width = Cm(15.5)
            shape.height = int(shape.height * ratio)
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.keep_with_next = True

    # Replace template header/footer content (including page-number content
    # controls) with a single centered page field.
    for part in (section.header, section.first_page_header):
        for child in list(part._element):
            part._element.remove(child)
        part.add_paragraph()
    footer = section.footer
    for child in list(footer._element):
        footer._element.remove(child)
    add_page_number(footer.add_paragraph())
    first_footer = section.first_page_footer
    for child in list(first_footer._element):
        first_footer._element.remove(child)
    first_footer.add_paragraph()

    doc.core_properties.title = "Dynamic Determination of Short-Term Corporate Credit Limits Using Machine Learning"
    doc.core_properties.subject = "Master's thesis in Computational Finance"
    doc.core_properties.author = "Đorđe Veličković"
    doc.core_properties.keywords = "corporate credit risk; credit limits; machine learning; SEC disclosures"
    doc.core_properties.comments = "Formatted from the RAF reference thesis."
    doc.save(path)


def build_thesis(page_map: dict[str, int], output: Path) -> None:
    source = BUILD_DIR / "combined_thesis.md"
    source.write_text(make_markdown(page_map), encoding="utf-8")
    run(
        "pandoc",
        str(source),
        "--from=markdown+pipe_tables+tex_math_dollars+raw_attribute",
        "--to=docx",
        f"--reference-doc={THESIS_TEMPLATE}",
        f"--resource-path={THESIS_DIR}:{REFERENCE_DIR}",
        "--wrap=none",
        f"--output={output}",
    )
    apply_thesis_format(output)


def render_pdf(docx_path: Path, pdf_path: Path) -> None:
    render_dir = BUILD_DIR / "rendered"
    render_dir.mkdir(parents=True, exist_ok=True)
    run(
        "soffice",
        "-env:UserInstallation=file:///private/tmp/wcl-lo-profile",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(render_dir),
        str(docx_path),
    )
    generated = render_dir / f"{docx_path.stem}.pdf"
    shutil.copy2(generated, pdf_path)


def extract_page_map(pdf_path: Path) -> tuple[dict[str, int], int]:
    headings = [title for _, title in parse_headings()]
    figures, tables = parse_captions()
    targets = headings + figures + tables
    pages = PdfReader(str(pdf_path)).pages
    page_texts = [(page.extract_text() or "").replace("\u00ad", "") for page in pages]
    mapping: dict[str, int] = {}
    for target in targets:
        compact_target = re.sub(r"\s+", " ", target).strip()
        # Pages 1–6 are the cover and front matter. Searching there would
        # mistake a TOC/list entry for the actual heading or caption.
        search_start = 1 if target == "Abstract" else 6
        for page_number, text in enumerate(page_texts[search_start:], start=search_start + 1):
            compact_text = re.sub(r"\s+", " ", text)
            caption_number = re.match(r"^(Figure|Table)\s+(\d+)\.", compact_target)
            # Caption numbers are globally unique. They are more reliable than
            # caption text because narrow tables can hyphenate the title.
            needle = caption_number.group(0) if caption_number else compact_target
            if needle in compact_text:
                mapping[target] = page_number
                break
    return mapping, len(pages)


def clear_document_paragraphs(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)


def configure_report_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(BASE_FONT_SIZE)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)


def add_report_paragraph(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = BASE_FONT_SIZE,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_indent: bool = True,
    before: float = 0,
    after: float = 6,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if first_indent and text:
        p.paragraph_format.first_line_indent = Cm(0.5)
    run = p.add_run(text)
    set_run_font(run, size, bold=bold, italic=italic)


def add_report_heading(doc: Document, number: int, title: str, *, page_break_before: bool = False) -> None:
    add_report_paragraph(
        doc,
        f"{number}. {title}",
        bold=True,
        size=16,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_indent=False,
        before=8,
        after=6,
    )
    if page_break_before:
        doc.paragraphs[-1].paragraph_format.page_break_before = True


def build_report(thesis_pages: int) -> None:
    doc = Document(REPORT_TEMPLATE)
    clear_document_paragraphs(doc)
    configure_report_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    add_report_paragraph(
        doc,
        "TO THE TEACHING AND SCIENTIFIC COUNCIL OF THE SCHOOL OF COMPUTING IN BELGRADE",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        after=18,
    )
    add_report_paragraph(
        doc,
        "The Teaching and Scientific Council of the School of Computing in Belgrade, at its [SESSION NUMBER] session held on [SESSION DATE], appointed us to the Commission for the review and evaluation of the master’s thesis submitted by Đorđe Veličković, student ID [STUDENT ID], under the title “Dynamic Determination of Short-Term Corporate Credit Limits Using Machine Learning”, prepared under the supervision of Prof. Dr. Branko Urošević. After reviewing the thesis, the Commission submits the following:",
        italic=True,
        after=18,
    )
    add_report_paragraph(
        doc,
        "REPORT",
        bold=True,
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        after=18,
    )

    add_report_heading(doc, 1, "Biographical data of the candidate")
    add_report_paragraph(
        doc,
        "Đorđe Veličković is a software and data engineer with professional experience in financial technology and large-scale data systems. He has worked at Rho as a Senior Data Engineer since 2022. Previously, he worked as a Software Engineer at Kindred Group plc / Hooloovoo from 2020 to 2022, as a Junior Software Engineer and Software Engineer at Centili from 2018 to 2020, and as a Software Engineer Intern at Infobip in 2017.",
    )
    add_report_paragraph(
        doc,
        "The candidate is enrolled in the Master in Computational Finance programme at the School of Computing, Union University. He completed a master’s programme in Software Engineering at the same institution in 2023 and a bachelor’s programme in Computer Science in 2019. His principal technical skills include SQL, Python and Java, with a professional focus on data engineering and software engineering. He has professional working proficiency in English.",
    )

    add_report_heading(doc, 2, "Description of the master thesis")
    add_report_paragraph(
        doc,
        f"The master’s thesis comprises {thesis_pages} pages in the submitted Word edition and cites 14 sources. It contains an abstract, eight chapters, a reference list and a reproducibility appendix. The thesis examines how point-in-time accounting information and machine-learning estimates can be combined in a transparent framework for dynamic short-term corporate credit limits.",
    )
    add_report_paragraph(
        doc,
        "The empirical study combines SEC Company Facts, Financial Statement and Notes data, and submissions in a monthly panel covering 2011–2025. The completed modeling panel contains 565,222 labeled company-month observations from 8,780 companies. Twelve accounting predictors describe liquidity, leverage, profitability, cash generation, growth and company size. The outcome is a new bankruptcy or reviewed payment-failure disclosure episode within 60 days; the candidate clearly distinguishes this disclosure proxy from verified default on a lending facility.",
    )
    add_report_paragraph(
        doc,
        "Four model families are compared using a chronological development design: logistic regression, L2-regularized logistic regression, random forest and XGBoost. The thesis then translates the selected probability into a preliminary limit through a financial-capacity proxy, a probability-based risk adjustment and an exposure cap. Monthly update rules introduce immediate reductions, controlled increases, a deadband and explicit hard stops.",
    )
    add_report_paragraph(
        doc,
        "The work is logically organized and supported by documented filters, formulas, tables, figures and reproducibility material. The candidate also states the main empirical limitations directly: the disclosure labels are automated and not independently validated, the later evaluation is retrospective, and facility utilization, repayment, loss, pricing and profitability data are unavailable.",
    )

    add_report_heading(doc, 3, "Analysis of the work and key results")
    add_report_paragraph(
        doc,
        "The selected random forest produces a meaningful ranking signal on 131,842 company-month observations from 2023–2025. It achieves ROC-AUC of 0.8866 and average precision of 2.5989%, or 10.77 times the 0.2412% event share. Its Brier score is 0.0023840, which is 0.95% better than the constant-probability benchmark. At the frozen alert threshold, precision is 5.63% and recall is 9.12%, so the results support cautious use as a risk-ranking input rather than a complete default-detection system.",
    )
    add_report_paragraph(
        doc,
        "The historical policy replay covers 99,619 observations from 3,869 companies. Under the reference assumptions, the median assigned limit is USD 4.23 million, the mean assigned limit is 0.668% of reported assets, and 72.05% of monthly transitions leave the limit unchanged. Sensitivity analysis shows that cash eligibility, assumed near-term debt payments and the exposure cap have a larger aggregate effect than the probability multiplier. This is an important practical result because it identifies the assumptions that require lender-specific evidence before implementation.",
    )

    add_report_heading(doc, 4, "Conclusion and proposal", page_break_before=True)
    add_report_paragraph(
        doc,
        "In the submitted master’s thesis, Đorđe Veličković demonstrates knowledge of corporate credit-risk modeling, chronological validation, rare-event evaluation and rule-based credit-limit design. The candidate connects the predictive and policy components coherently, documents the empirical evidence, and interprets the results with appropriate caution.",
    )
    add_report_paragraph(
        doc,
        "Subject to confirmation of the official administrative details and the signatures below, the Commission proposes that the Teaching and Scientific Council accept the master’s thesis entitled “Dynamic Determination of Short-Term Corporate Credit Limits Using Machine Learning” and approve its public oral defense.",
        after=18,
    )

    add_report_paragraph(
        doc,
        "Belgrade, [REPORT DATE]\t\t\tMembers of the Commission:",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_indent=False,
        after=24,
    )
    add_report_paragraph(
        doc,
        "_________________________________________\nProf. Dr. Branko Urošević ([ROLE])",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_indent=False,
        after=24,
    )
    add_report_paragraph(
        doc,
        "_________________________________________\n[COMMISSION MEMBER] ([ROLE])",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_indent=False,
        after=24,
    )
    add_report_paragraph(
        doc,
        "_________________________________________\n[COMMISSION MEMBER] ([ROLE])",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_indent=False,
    )

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text.strip():
                            set_run_font(run, BASE_FONT_SIZE, bold=run.bold)

    doc.core_properties.title = "Report on the completed master's thesis of Đorđe Veličković"
    doc.core_properties.subject = "Draft report for the Teaching and Scientific Council"
    doc.core_properties.author = "Commission draft"
    doc.core_properties.comments = "Administrative placeholders must be completed and approved by the Commission."
    doc.save(REPORT_DOCX)


def validate_docx(path: Path, forbidden: tuple[str, ...] = ()) -> None:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    for value in forbidden:
        if value.casefold() in text.casefold():
            raise RuntimeError(f"Unexpected template text {value!r} remains in {path.name}")


def build_thesis_word_only() -> dict[str, object]:
    """Build only the thesis DOCX, using temporary renders for page numbers."""
    first_docx = BUILD_DIR / "word_only_first_pass.docx"
    first_pdf = BUILD_DIR / "word_only_first_pass.pdf"
    validation_pdf = BUILD_DIR / "word_only_pagination.pdf"

    build_thesis({}, first_docx)
    render_pdf(first_docx, first_pdf)
    page_map, _ = extract_page_map(first_pdf)

    thesis_pages = 0
    final_map = page_map
    for _ in range(3):
        build_thesis(page_map, THESIS_DOCX)
        render_pdf(THESIS_DOCX, validation_pdf)
        final_map, thesis_pages = extract_page_map(validation_pdf)
        if final_map == page_map:
            break
        page_map = final_map
    else:
        build_thesis(final_map, THESIS_DOCX)

    validate_docx(THESIS_DOCX, ("Petar Damnjanović", "Identifying Financial Market Regimes"))
    targets = [title for _, title in parse_headings()] + list(sum(parse_captions(), []))
    return {
        "thesis_pages": thesis_pages,
        "toc_entries": len(parse_headings()),
        "figures": len(parse_captions()[0]),
        "tables": len(parse_captions()[1]),
        "missing_page_targets": [title for title in targets if title not in final_map],
        "outputs": [str(THESIS_DOCX)],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--word-only",
        action="store_true",
        help="build only the thesis DOCX; temporary PDFs are used only to calculate page numbers",
    )
    args = parser.parse_args()
    if not THESIS_TEMPLATE.exists():
        raise SystemExit(f"Thesis reference document not found: {THESIS_TEMPLATE}")
    if not args.word_only and not REPORT_TEMPLATE.exists():
        raise SystemExit(f"Report reference document not found: {REPORT_TEMPLATE}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    BUILD_DIR.mkdir(parents=True, exist_ok=True)

    if args.word_only:
        print(json.dumps(build_thesis_word_only(), indent=2, ensure_ascii=False))
        return

    first_docx = BUILD_DIR / "thesis_first_pass.docx"
    first_pdf = BUILD_DIR / "thesis_first_pass.pdf"
    build_thesis({}, first_docx)
    render_pdf(first_docx, first_pdf)
    page_map, _ = extract_page_map(first_pdf)

    build_thesis(page_map, THESIS_DOCX)
    render_pdf(THESIS_DOCX, THESIS_PDF)
    final_map, thesis_pages = extract_page_map(THESIS_PDF)
    if final_map != page_map:
        build_thesis(final_map, THESIS_DOCX)
        render_pdf(THESIS_DOCX, THESIS_PDF)
        final_map, thesis_pages = extract_page_map(THESIS_PDF)

    build_report(thesis_pages)
    render_pdf(REPORT_DOCX, REPORT_PDF)
    validate_docx(THESIS_DOCX, ("Petar Damnjanović", "Identifying Financial Market Regimes"))
    validate_docx(REPORT_DOCX, ("Petar Damnjanović", "Mikica Drenovak", "Vladimir Vasić"))

    missing = [title for title in [t for _, t in parse_headings()] + list(sum(parse_captions(), [])) if title not in final_map]
    summary = {
        "thesis_pages": thesis_pages,
        "report_pages": len(PdfReader(str(REPORT_PDF)).pages),
        "toc_entries": len(parse_headings()),
        "figures": len(parse_captions()[0]),
        "tables": len(parse_captions()[1]),
        "missing_page_targets": missing,
        "outputs": [str(THESIS_DOCX), str(THESIS_PDF), str(REPORT_DOCX), str(REPORT_PDF)],
    }
    (OUTPUT_DIR / "build_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
